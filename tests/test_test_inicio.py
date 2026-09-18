from io import BytesIO
from docx import Document
from fastapi.testclient import TestClient
from main import app, _firmar_sesion, _parsear_datos_academicos
from test_module.web import _AUTOMATION_JOBS, _AUTOMATION_SYNC, _JOBS
import threading


def _session_client() -> TestClient:
    client = TestClient(app)
    client.cookies.set("ap_sesion", _firmar_sesion("adminlocal"))
    return client


def _crear_docx_prueba() -> bytes:
    doc = Document()
    doc.add_heading("EVALUACIÓN TEST", level=1)
    p = doc.add_paragraph()
    p.add_run("1. ¿Cuál es el objeto de la macroeconomía?\n")
    p.add_run("A) Los agregados económicos\n")
    p.add_run("B) La contabilidad individual\n")
    p.add_run("Respuesta correcta: A\n")
    b = BytesIO()
    doc.save(b)
    return b.getvalue()


def test_parsear_datos_academicos_espacios():
    res = _parsear_datos_academicos("ADM101 PEN2024 MAC001 Macroeconomía Aplicada")
    assert res is not None
    assert res["programa"] == "ADM101"
    assert res["pensum"] == "PEN2024"
    assert res["codigo_asignatura"] == "MAC001"
    assert res["asignatura"] == "Macroeconomía Aplicada"


def test_parsear_datos_academicos_guiones():
    res = _parsear_datos_academicos("ADM101-PEN2024-MAC001-Macroeconomía Aplicada")
    assert res is not None
    assert res["programa"] == "ADM101"
    assert res["pensum"] == "PEN2024"
    assert res["codigo_asignatura"] == "MAC001"
    assert res["asignatura"] == "Macroeconomía Aplicada"


def test_parsear_datos_academicos_guiones_espacios():
    res = _parsear_datos_academicos("ADM101 - PEN2024 - MAC001 - Macroeconomía Aplicada")
    assert res is not None
    assert res["programa"] == "ADM101"
    assert res["pensum"] == "PEN2024"
    assert res["codigo_asignatura"] == "MAC001"
    assert res["asignatura"] == "Macroeconomía Aplicada"


def test_parsear_datos_academicos_invalido():
    assert _parsear_datos_academicos("ADM101 PEN2024") is None
    assert _parsear_datos_academicos("") is None
    assert _parsear_datos_academicos("   ") is None


def test_pagina_inicio_contiene_test_evaluativo():
    client = _session_client()
    response = client.get("/")
    assert response.status_code == 200
    html = response.text
    # Debe contener la tarjeta de Test Evaluativo
    assert 'name="test_evaluativo"' in html
    # Debe estar desmarcado por defecto (no tener 'checked' en el input de test)
    assert '<input type="checkbox" name="test_evaluativo" value="1" id="chk-test">' in html
    # Debe contener el bloque de test
    assert 'id="bloque-test"' in html
    assert 'id="test-datos-academicos"' in html
    assert 'id="test-semestre"' in html
    assert 'id="archivo-test-docx"' in html
    assert 'id="test-google-doc-url"' in html


def test_analizar_solo_test_flujo_exclusivo():
    client = _session_client()
    doc_bytes = _crear_docx_prueba()

    # Enviar solo test sin OVI, OVA ni Retos, y sin matriz ni zip
    response = client.post(
        "/analizar",
        data={
            "test_evaluativo": "1",
            "test_datos_academicos": "ECO01 PEN2024 MAC101 Macroeconomía I",
            "test_semestre": "3",
            "modo_test": "archivo",
        },
        files={
            "test_archivo_docx": ("test_eval.docx", BytesIO(doc_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        },
        follow_redirects=False,
    )

    # Debe redirigir a /tests/{job_id}/revision
    assert response.status_code == 303
    location = response.headers.get("location", "")
    assert location.startswith("/tests/")
    assert location.endswith("/revision")

    # Verificar que los datos académicos quedaron registrados
    job_id = location.split("/")[2]
    assert job_id in _JOBS
    cascada = _JOBS[job_id]["cascada_datos"]
    assert cascada["programa"] == "ECO01"
    assert cascada["pensum"] == "PEN2024"
    assert cascada["codigo_asignatura"] == "MAC101"
    assert cascada["asignatura"] == "Macroeconomía I"
    assert cascada["nivel"] == "3"


def test_seleccionar_asignatura_endpoint():
    client = _session_client()
    job_id = "test_job_sync_123"
    _JOBS[job_id] = {"owner": "adminlocal", "document": None}
    sync_obj = {"event": threading.Event(), "seleccion": {"asignatura": None}}
    _AUTOMATION_SYNC[job_id] = sync_obj
    _AUTOMATION_JOBS[job_id] = {
        "owner": "adminlocal",
        "estado": "esperando_asignatura",
        "esperando_asignatura": True,
        "opciones_asignatura": ["Macroeconomía I (MAC101)", "Macroeconomía II (MAC102)"],
        "asignatura_buscada": "Macroeconomía",
    }

    resp = client.post(
        f"/tests/{job_id}/seleccionar-asignatura",
        data={"asignatura": "Macroeconomía I (MAC101)"},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["ok"] is True
    assert data["asignatura"] == "Macroeconomía I (MAC101)"
    assert sync_obj["seleccion"]["asignatura"] == "Macroeconomía I (MAC101)"
    assert sync_obj["event"].is_set()


def test_facultad_opcional_en_automatizacion_y_revision():
    from unittest.mock import MagicMock
    from test_module.automation import PrizmaTestAutomator

    # 1. Verificar que seleccionar_mui con Facultad retorna True si el selector no existe
    automator = PrizmaTestAutomator(base_url="https://admin.prizma.solutions", headless=True)
    mock_page = MagicMock()
    # Simular que no se localiza ningún campo
    mock_page.locator.return_value.count.return_value = 0
    mock_page.get_by_placeholder.return_value.count.return_value = 0
    mock_page.get_by_label.return_value.count.return_value = 0

    # Para cualquier campo obligatorio, debe lanzar ValueError
    try:
        automator.seleccionar_mui(mock_page, "Programa", "ADM101")
        assert False, "Debió fallar para Programa no encontrado"
    except ValueError:
        pass

    # Para Facultad, al ser opcional, debe retornar True sin lanzar excepción
    resultado_fac = automator.seleccionar_mui(mock_page, "Facultad", "Virtual - Derecho")
    assert resultado_fac is True

    # 2. Verificar que en la página de revisión el campo Facultad esté etiquetado como Opcional
    client = _session_client()
    doc_bytes = _crear_docx_prueba()
    resp_upload = client.post(
        "/tests/analizar",
        files={"archivo_docx": ("test.docx", BytesIO(doc_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"modo_entrada": "archivo"},
        follow_redirects=False,
    )
    job_id = resp_upload.headers.get("location", "").split("/")[-2]
    resp_rev = client.get(f"/tests/{job_id}/revision")
    assert resp_rev.status_code == 200
    assert "Facultad" in resp_rev.text
    assert "Opcional" in resp_rev.text

