from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from main import app, _firmar_sesion
from test_module.models import document_to_dict
from test_module.parser import parse_docx_bytes


def _session_client() -> TestClient:
    client = TestClient(app)
    client.cookies.set("ap_sesion", _firmar_sesion("adminlocal"))
    return client


def _docx_fixture() -> bytes:
    document = Document()
    document.add_paragraph("Título del TEST")
    document.add_paragraph("Prueba de interfaz")
    document.add_paragraph("Semana 1 - Corte 1")
    document.add_paragraph("Verdadero/Falso")
    document.add_paragraph("1. El parser funciona. [R: Verdadero]")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_tests_page_is_available_to_authenticated_user():
    response = _session_client().get("/tests")

    assert response.status_code == 200
    assert "Nuevo Test" in response.text
    assert "Asistente" in response.text


def test_analyze_docx_redirects_to_review_with_job_id():
    response = _session_client().post(
        "/tests/analizar",
        files={"archivo_docx": ("prueba.docx", _docx_fixture(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"].startswith("/tests/")
    assert response.headers["location"].endswith("/revision")


def test_review_and_download_expose_parsed_test_without_prizma_side_effects():
    client = _session_client()
    analyzed = client.post(
        "/tests/analizar",
        files={"archivo_docx": ("prueba.docx", _docx_fixture(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        follow_redirects=False,
    )
    review_url = analyzed.headers["location"]

    review = client.get(review_url)
    download = client.get(review_url.replace("/revision", "/descargar"))

    assert review.status_code == 200
    assert "Prueba de interfaz" in review.text
    assert "Verdadero o falso" in review.text
    assert download.status_code == 200
    assert download.headers["content-type"].startswith("application/json")
    assert b"Prueba de interfaz" in download.content


def test_analyze_google_doc_uses_configured_loader():
    app.state.test_google_loader = lambda url, user: (_docx_fixture(), "google.docx")
    try:
        response = _session_client().post(
            "/tests/analizar",
            data={"google_doc_url": "https://docs.google.com/document/d/example/edit"},
            follow_redirects=False,
        )
    finally:
        del app.state.test_google_loader

    assert response.status_code == 303
    assert response.headers["location"].endswith("/revision")


def test_save_review_persists_edited_title_and_prompt():
    client = _session_client()
    analyzed = client.post(
        "/tests/analizar",
        files={"archivo_docx": ("prueba.docx", _docx_fixture(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        follow_redirects=False,
    )
    job_id = analyzed.headers["location"].split("/")[2]
    state = document_to_dict(parse_docx_bytes(_docx_fixture()))
    state["tests"][0]["title"] = "Título corregido"
    state["tests"][0]["sections"][0]["questions"][0]["prompt"] = "Pregunta corregida."

    saved = client.post(f"/tests/{job_id}/guardar", data={"payload": __import__("json").dumps(state)})
    download = client.get(f"/tests/{job_id}/descargar")

    assert saved.status_code == 200
    assert "Título corregido" in saved.text
    assert "Título corregido".encode("utf-8") in download.content
    assert b"Pregunta corregida." in download.content


def test_tests_page_includes_dual_mode_and_help_panel():
    response = _session_client().get("/tests")
    assert response.status_code == 200
    assert "Por archivo" in response.text
    assert "Por link" in response.text
    assert "¿Cómo funciona?" in response.text
    assert "dropzone" in response.text


def test_cargue_actual_retains_test_nav_item():
    response = _session_client().get("/cargue-actual")
    assert response.status_code == 200
    assert 'href="/tests"' in response.text


def test_review_page_highlights_erroneous_question_with_alert_and_link():
    # Docx fixture with a question that has no answer
    source = Document()
    source.add_paragraph("Título del TEST")
    source.add_paragraph("Examen con Errores")
    source.add_paragraph("Verdadero/Falso")
    source.add_paragraph("1. Pregunta sin respuesta.")
    buffer = BytesIO()
    source.save(buffer)

    client = _session_client()
    analyzed = client.post(
        "/tests/analizar",
        data={"modo_entrada": "archivo"},
        files={"archivo_docx": ("error.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        follow_redirects=False,
    )
    assert analyzed.status_code == 303
    review = client.get(analyzed.headers["location"])
    assert review.status_code == 200
    assert "question-error" in review.text
    assert "Corrección requerida en Pregunta 1" in review.text
    assert "href=\"#pregunta-0-0-0\"" in review.text
    assert "Requiere atención" in review.text

