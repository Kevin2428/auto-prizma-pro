from io import BytesIO
from unittest.mock import MagicMock, patch

from docx import Document
from fastapi.testclient import TestClient

from main import app, _firmar_sesion
from test_module.automation import PrizmaTestAutomator, normalizar_texto
from test_module.models import DocumentTest, Question, Section


def _client() -> TestClient:
    client = TestClient(app)
    client.cookies.set("ap_sesion", _firmar_sesion("adminlocal"))
    return client


def test_normalizar_texto_removes_accents_and_lowercases():
    assert normalizar_texto("  ÁÉÍÓÚ ñ Ñ  ") == "aeiou n n"
    assert normalizar_texto("Preguntas de Única Respuesta") == "preguntas de unica respuesta"


def test_automator_initialization_respects_headful_and_url():
    auto = PrizmaTestAutomator(base_url="https://admin.prizma.solutions/", headless=False)
    assert auto.base_url == "https://admin.prizma.solutions"
    assert auto.headless is False


def test_review_page_renders_phase3_prizma_panel():
    client = _client()
    source = Document()
    source.add_paragraph("Título del TEST")
    source.add_paragraph("Examen Fase 3")
    source.add_paragraph("Verdadero/Falso")
    source.add_paragraph("1. Prueba automatizada. [R: Verdadero]")
    buf = BytesIO()
    source.save(buf)

    analyzed = client.post(
        "/tests/analizar",
        data={"modo_entrada": "archivo"},
        files={"archivo_docx": ("test_fase3.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        follow_redirects=False,
    )
    assert analyzed.status_code == 303
    job_id = analyzed.headers["location"].split("/")[2]

    review = client.get(f"/tests/{job_id}/revision")
    assert review.status_code == 200
    assert "Fase 3: Publicación y Carga en PRIZMA" in review.text
    assert "prizma_usuario" in review.text
    assert "prizma_clave" in review.text
    assert "bloque-acceso-prizma" in review.text
    assert "cascada_facultad" in review.text
    assert "cascada_asignatura" in review.text
    assert "btn-lanzar-prizma" in review.text
    assert "Crear y Publicar Evaluación en PRIZMA" in review.text
    assert "Ir a Carga en PRIZMA" in review.text


def test_ejecutar_prizma_blocks_if_errors_exist():
    client = _client()
    source = Document()
    source.add_paragraph("Título del TEST")
    source.add_paragraph("Examen con Fallos")
    source.add_paragraph("Verdadero/Falso")
    source.add_paragraph("1. Pregunta sin clave.")
    buf = BytesIO()
    source.save(buf)

    analyzed = client.post(
        "/tests/analizar",
        data={"modo_entrada": "archivo"},
        files={"archivo_docx": ("fallos.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        follow_redirects=False,
    )
    job_id = analyzed.headers["location"].split("/")[2]

    resp = client.post(
        f"/tests/{job_id}/ejecutar-prizma",
        data={
            "prizma_usuario": "admin",
            "prizma_clave": "12345",
        },
    )
    assert resp.status_code == 400
    assert "Corrige los errores pendientes" in resp.text


def test_estado_prizma_monitoring_and_json():
    client = _client()
    source = Document()
    source.add_paragraph("Título del TEST")
    source.add_paragraph("Examen Valido")
    source.add_paragraph("Verdadero/Falso")
    source.add_paragraph("1. Todo ok. [R: Verdadero]")
    buf = BytesIO()
    source.save(buf)

    analyzed = client.post(
        "/tests/analizar",
        data={"modo_entrada": "archivo"},
        files={"archivo_docx": ("valido.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        follow_redirects=False,
    )
    job_id = analyzed.headers["location"].split("/")[2]

    # Mock PrizmaTestAutomator.ejecutar_cargue so test doesn't launch real browser
    with patch("test_module.web.PrizmaTestAutomator") as mock_cls:
        instance = mock_cls.return_value
        instance.ejecutar_cargue.return_value = {"ok": True, "simulacion": True, "procesados": 1, "total": 1}

        resp = client.post(
            f"/tests/{job_id}/ejecutar-prizma",
            data={
                "prizma_usuario": "usuario_test",
                "prizma_clave": "clave_test",
                "recordar_credenciales": "1",
                "solo_simulacion": "1",
            },
            follow_redirects=False,
        )
        assert resp.status_code == 303
        assert resp.headers["location"] == f"/tests/{job_id}/estado-prizma"

        # Check monitoring page
        mon = client.get(f"/tests/{job_id}/estado-prizma")
        assert mon.status_code == 200
        assert "Monitoreo de Simulación en PRIZMA" in mon.text
        assert "tabla-tests-body" in mon.text

        # Check json endpoint
        mon_json = client.get(f"/tests/{job_id}/estado-prizma/json")
        assert mon_json.status_code == 200
        data = mon_json.json()
        assert "estado" in data
        assert data["total"] == 1
        assert data.get("simulacion") is True


def test_seleccionar_mui_raises_error_if_option_not_found():
    import pytest
    automator = PrizmaTestAutomator()
    automator._tomar_captura = MagicMock(return_value="")

    mock_page = MagicMock()
    mock_input = MagicMock()
    mock_input.count.return_value = 1
    mock_input.first = mock_input
    mock_input.input_value.return_value = ""
    mock_page.locator.return_value = mock_input

    # Emulate option list with values that do NOT match
    mock_option = MagicMock()
    mock_option.is_visible.return_value = True
    mock_option.inner_text.return_value = "Pensum Antiguo 2010"

    mock_options_loc = MagicMock()
    mock_options_loc.count.return_value = 1
    mock_options_loc.nth.return_value = mock_option

    def locator_side_effect(selector):
        if "option" in selector:
            return mock_options_loc
        return mock_input

    mock_page.locator.side_effect = locator_side_effect

    with pytest.raises(ValueError, match="La opción o código 'Pensum Inexistente 2099' no existe"):
        automator.seleccionar_mui(mock_page, "Pensum", "Pensum Inexistente 2099")


def test_seleccionar_mui_matches_by_code():
    automator = PrizmaTestAutomator()
    automator._tomar_captura = MagicMock(return_value="")

    mock_page = MagicMock()
    mock_input = MagicMock()
    mock_input.count.return_value = 1
    mock_input.first = mock_input
    mock_input.input_value.return_value = ""

    mock_option = MagicMock()
    mock_option.is_visible.return_value = True
    mock_option.inner_text.return_value = "COD111 - NEGOCIOS INTERNACIONALES"

    mock_options_loc = MagicMock()
    mock_options_loc.count.return_value = 1
    mock_options_loc.nth.return_value = mock_option

    def locator_side_effect(selector):
        if "option" in selector:
            return mock_options_loc
        return mock_input

    mock_page.locator.side_effect = locator_side_effect

    # Passing only the code "COD111" should match the option containing that code
    assert automator.seleccionar_mui(mock_page, "Programa", "COD111") is True
    mock_option.click.assert_called_once()


def test_seleccionar_mui_detects_ambiguity_when_code_missing():
    import pytest
    automator = PrizmaTestAutomator()
    automator._tomar_captura = MagicMock(return_value="")

    mock_page = MagicMock()
    mock_input = MagicMock()
    mock_input.count.return_value = 1
    mock_input.first = mock_input
    mock_input.input_value.return_value = ""

    # Two options with same name but different codes
    mock_opt1 = MagicMock()
    mock_opt1.is_visible.return_value = True
    mock_opt1.inner_text.return_value = "COD101 - DERECHO"

    mock_opt2 = MagicMock()
    mock_opt2.is_visible.return_value = True
    mock_opt2.inner_text.return_value = "COD102 - DERECHO"

    mock_options_loc = MagicMock()
    mock_options_loc.count.return_value = 2
    mock_options_loc.nth.side_effect = [mock_opt1, mock_opt2, mock_opt1, mock_opt2]

    def locator_side_effect(selector):
        if "option" in selector:
            return mock_options_loc
        return mock_input

    mock_page.locator.side_effect = locator_side_effect

    # Searching with generic name "DERECHO" without code should halt with ambiguity error
    with pytest.raises(ValueError, match="Ambigüedad en 'Programa'.*ingresa el código exacto"):
        automator.seleccionar_mui(mock_page, "Programa", "DERECHO")


def test_navegar_a_crear_test_clicks_actividades_then_test():
    automator = PrizmaTestAutomator()
    automator._tomar_captura = MagicMock(return_value="")

    mock_page = MagicMock()

    mock_btn_act = MagicMock()
    mock_btn_act.count.return_value = 1
    mock_btn_act.first = mock_btn_act
    mock_btn_act.is_visible.return_value = True

    mock_btn_test = MagicMock()
    mock_btn_test.count.return_value = 1
    mock_btn_test.first = mock_btn_test
    mock_btn_test.is_visible.return_value = True

    mock_btn_crear = MagicMock()
    mock_btn_crear.count.return_value = 1
    mock_btn_crear.first = mock_btn_crear
    mock_btn_crear.is_visible.return_value = True

    form_visible = False

    def on_crear_click(*args, **kwargs):
        nonlocal form_visible
        form_visible = True

    mock_btn_crear.click.side_effect = on_crear_click

    mock_campo_fac = MagicMock()
    mock_campo_fac.count.return_value = 1
    mock_campo_fac.first = mock_campo_fac
    mock_campo_fac.is_visible.side_effect = lambda: form_visible

    def loc_side_effect(sel):
        if "Actividades" in sel:
            return mock_btn_act
        if "Test" in sel:
            return mock_btn_test
        if "Crear" in sel:
            return mock_btn_crear
        if "Facultad" in sel:
            return mock_campo_fac
        elem = MagicMock()
        elem.count.return_value = 0
        elem.is_visible.return_value = False
        elem.first = elem
        return elem

    mock_page.locator.side_effect = loc_side_effect

    mock_btn_cancelar = MagicMock()
    mock_btn_cancelar.count.return_value = 0
    mock_btn_cancelar.first = mock_btn_cancelar
    mock_btn_cancelar.is_visible.return_value = False

    def role_effect(role, name=None, **kwargs):
        name_str = str(name or "")
        if "cancelar" in name_str.lower():
            return mock_btn_cancelar
        if "crear" in name_str.lower() or "crear" in str(role).lower():
            return mock_btn_crear
        if "actividades" in name_str.lower():
            return mock_btn_act
        if "test" in name_str.lower():
            return mock_btn_test
        m = MagicMock()
        m.first = m
        m.is_visible.return_value = False
        return m

    mock_page.get_by_role.side_effect = role_effect

    automator.navegar_a_crear_test(mock_page)

    assert mock_page.get_by_role.call_count >= 2
    assert mock_btn_crear.click.call_count >= 1


def test_localizar_campo_formulario_finds_visible_input():
    automator = PrizmaTestAutomator()
    mock_page = MagicMock()

    mock_visible_input = MagicMock()
    mock_visible_input.count.return_value = 1
    mock_visible_input.first = mock_visible_input
    mock_visible_input.is_visible.return_value = True

    mock_page.get_by_label.return_value.count.return_value = 1
    mock_page.get_by_label.return_value.nth.return_value = mock_visible_input

    res = automator.localizar_campo_formulario(mock_page, "Facultad", timeout_ms=500)
    assert res == mock_visible_input


def test_ejecutar_cargue_simulation_mode_does_not_click_guardar():
    automator = PrizmaTestAutomator(headless=True)
    automator._tomar_captura = MagicMock(return_value="")
    automator.login = MagicMock(return_value=True)
    automator.navegar_a_crear_test = MagicMock()
    automator.seleccionar_mui = MagicMock(return_value=True)
    automator.inyectar_json_test = MagicMock(return_value=True)

    test = DocumentTest(title="Test Simulado", sections=[])

    with patch("test_module.automation.sync_playwright") as mock_sp:
        mock_p = MagicMock()
        mock_sp.return_value.__enter__.return_value = mock_p

        mock_page = MagicMock()
        mock_p.chromium.launch.return_value.new_context.return_value.new_page.return_value = mock_page

        mock_btn_guardar = MagicMock()
        mock_btn_guardar.first = mock_btn_guardar

        def loc_effect(sel):
            if "Guardar" in sel:
                return mock_btn_guardar
            mock_elem = MagicMock()
            mock_elem.first = mock_elem
            return mock_elem

        mock_page.locator.side_effect = loc_effect

        # Execute in simulation mode
        res = automator.ejecutar_cargue(
            usuario_prizma="usuario",
            clave_prizma="clave",
            tests=[test],
            cascada={},
            solo_simulacion=True,
        )

        assert res["ok"] is True
        assert res["simulacion"] is True
        assert res["procesados"] == 1
        # In simulation mode, btn_guardar.click must NEVER be called
        mock_btn_guardar.click.assert_not_called()


def test_resolver_variantes_numericas():
    from test_module.automation import resolver_variantes_numericas

    vars_3 = resolver_variantes_numericas("3")
    assert "tres" in vars_3
    assert "tercero" in vars_3
    assert "iii" in vars_3

    vars_word = resolver_variantes_numericas("Tres")
    assert "3" in vars_word
    assert "tres" in vars_word

    vars_level = resolver_variantes_numericas("Nivel 3")
    assert "3" in vars_level
    assert "tres" in vars_level

    vars_non_num = resolver_variantes_numericas("Derecho")
    assert vars_non_num == []


def test_seleccionar_mui_numeric_bidirectional_digit_to_word():
    automator = PrizmaTestAutomator()
    automator._tomar_captura = MagicMock(return_value="")

    mock_input = MagicMock()
    mock_input.is_visible.return_value = True
    mock_input.input_value.return_value = ""

    automator.localizar_campo_formulario = MagicMock(return_value=mock_input)

    mock_page = MagicMock()

    # Option in PRIZMA is written in words: "Tres"
    mock_option = MagicMock()
    mock_option.is_visible.return_value = True
    mock_option.inner_text.return_value = "Tres"

    mock_options_loc = MagicMock()
    mock_options_loc.count.return_value = 1
    mock_options_loc.nth.return_value = mock_option

    mock_page.locator.return_value = mock_options_loc

    # User inputs digit "3" for "Nivel de pensum", should match "Tres"
    assert automator.seleccionar_mui(mock_page, "Nivel de pensum", "3") is True
    mock_option.click.assert_called_once()


def test_seleccionar_mui_numeric_bidirectional_word_to_digit():
    automator = PrizmaTestAutomator()
    automator._tomar_captura = MagicMock(return_value="")

    mock_input = MagicMock()
    mock_input.is_visible.return_value = True
    mock_input.input_value.return_value = ""

    automator.localizar_campo_formulario = MagicMock(return_value=mock_input)

    mock_page = MagicMock()

    # Option in PRIZMA is written with digits: "Nivel 3"
    mock_option = MagicMock()
    mock_option.is_visible.return_value = True
    mock_option.inner_text.return_value = "Nivel 3"

    mock_options_loc = MagicMock()
    mock_options_loc.count.return_value = 1
    mock_options_loc.nth.return_value = mock_option

    mock_page.locator.return_value = mock_options_loc

    # User inputs word "Tercero" or "Tres", should match "Nivel 3"
    assert automator.seleccionar_mui(mock_page, "Nivel de pensum", "Tercero") is True
    mock_option.click.assert_called_once()


def test_llenar_descripcion_test_contenteditable():
    automator = PrizmaTestAutomator()
    automator._tomar_captura = MagicMock(return_value="")

    mock_page = MagicMock()
    mock_editor = MagicMock()
    mock_editor.is_visible.return_value = True

    mock_loc = MagicMock()
    mock_loc.count.return_value = 1
    mock_loc.nth.return_value = mock_editor

    mock_page.locator.return_value = mock_loc

    res = automator.llenar_descripcion_test(mock_page, "Esta es una descripción de prueba.")
    assert res is True
    mock_editor.click.assert_called_once()
    mock_editor.fill.assert_called_once_with("Esta es una descripción de prueba.")


def test_llenar_descripcion_test_contenteditable_keyboard_fallback():
    automator = PrizmaTestAutomator()
    automator._tomar_captura = MagicMock(return_value="")

    mock_page = MagicMock()
    mock_editor = MagicMock()
    mock_editor.is_visible.return_value = True
    # If .fill() is not supported on pure contenteditable div, it raises an exception and falls back to keyboard
    mock_editor.fill.side_effect = Exception("Element is not an <input> or <textarea>")

    mock_loc = MagicMock()
    mock_loc.count.return_value = 1
    mock_loc.nth.return_value = mock_editor

    mock_page.locator.return_value = mock_loc

    res = automator.llenar_descripcion_test(mock_page, "Descripción enriquecida.")
    assert res is True
    mock_editor.press.assert_any_call("Control+A")
    mock_editor.press.assert_any_call("Backspace")
    mock_page.keyboard.insert_text.assert_called_once_with("Descripción enriquecida.")


def test_navegar_a_crear_test_forzar_nuevo_cancels_previous():
    automator = PrizmaTestAutomator()
    automator._tomar_captura = MagicMock(return_value="")

    mock_page = MagicMock()

    mock_btn_cancelar = MagicMock()
    mock_btn_cancelar.count.return_value = 1
    mock_btn_cancelar.first = mock_btn_cancelar
    mock_btn_cancelar.is_visible.return_value = True

    mock_btn_act = MagicMock()
    mock_btn_act.count.return_value = 1
    mock_btn_act.first = mock_btn_act
    mock_btn_act.is_visible.return_value = True

    mock_btn_test = MagicMock()
    mock_btn_test.count.return_value = 1
    mock_btn_test.first = mock_btn_test
    mock_btn_test.is_visible.return_value = True

    mock_btn_crear = MagicMock()
    mock_btn_crear.count.return_value = 1
    mock_btn_crear.first = mock_btn_crear
    mock_btn_crear.is_visible.return_value = True

    mock_campo_fac = MagicMock()
    mock_campo_fac.count.return_value = 1
    mock_campo_fac.first = mock_campo_fac
    mock_campo_fac.is_visible.return_value = True

    def loc_side_effect(sel):
        if "Actividades" in sel:
            return mock_btn_act
        if "Test" in sel:
            return mock_btn_test
        if "Crear" in sel:
            return mock_btn_crear
        if "Facultad" in sel:
            return mock_campo_fac
        elem = MagicMock()
        elem.count.return_value = 0
        elem.is_visible.return_value = False
        elem.first = elem
        return elem

    mock_page.locator.side_effect = loc_side_effect

    def role_effect(role, name=None, **kwargs):
        name_str = str(name or "")
        if "cancelar" in name_str.lower():
            return mock_btn_cancelar
        if "crear" in name_str.lower() or "crear" in str(role).lower():
            return mock_btn_crear
        if "actividades" in name_str.lower():
            return mock_btn_act
        if "test" in name_str.lower():
            return mock_btn_test
        m = MagicMock()
        m.first = m
        m.is_visible.return_value = False
        return m

    mock_page.get_by_role.side_effect = role_effect

    automator.navegar_a_crear_test(mock_page, forzar_nuevo=True)

    mock_btn_cancelar.click.assert_called_once()
    assert mock_btn_crear.click.call_count >= 1


def test_guardar_creacion_test_locates_button_and_confirms_save():
    automator = PrizmaTestAutomator()
    automator._tomar_captura = MagicMock(return_value="")

    mock_page = MagicMock()

    mock_btn_guardar = MagicMock()
    mock_btn_guardar.is_visible.return_value = True
    mock_btn_guardar.count.return_value = 1
    mock_btn_guardar.inner_text.return_value = "Guardar"

    mock_hermano_loc = MagicMock()
    mock_hermano_loc.count.return_value = 1
    mock_hermano_loc.first = mock_btn_guardar

    mock_campo_nombre = MagicMock()
    mock_campo_nombre.count.return_value = 0
    mock_campo_nombre.is_visible.return_value = False

    def loc_side_effect(sel):
        if "Cancelar" in sel:
            return mock_hermano_loc
        if "Nombre del test" in sel:
            return mock_campo_nombre
        elem = MagicMock()
        elem.count.return_value = 0
        elem.is_visible.return_value = False
        elem.first = elem
        return elem

    mock_page.locator.side_effect = loc_side_effect

    res = automator.guardar_creacion_test(mock_page, "Examen Final de Prueba")
    assert res is True
    mock_btn_guardar.click.assert_called_once()


def test_guardar_creacion_test_raises_when_prizma_shows_error_alert():
    import pytest
    automator = PrizmaTestAutomator()
    automator._tomar_captura = MagicMock(return_value="")

    mock_page = MagicMock()

    mock_btn_guardar = MagicMock()
    mock_btn_guardar.is_visible.return_value = True
    mock_btn_guardar.count.return_value = 1
    mock_btn_guardar.inner_text.return_value = "Crear"

    mock_hermano_loc = MagicMock()
    mock_hermano_loc.count.return_value = 1
    mock_hermano_loc.first = mock_btn_guardar

    mock_error_alert = MagicMock()
    mock_error_alert.count.return_value = 1
    mock_error_alert.first = mock_error_alert
    mock_error_alert.is_visible.return_value = True
    mock_error_alert.inner_text.return_value = "El campo código es obligatorio."

    def loc_side_effect(sel):
        if "Cancelar" in sel:
            return mock_hermano_loc
        if "MuiAlert-standardError" in sel or "role=\"alert\"" in sel:
            return mock_error_alert
        elem = MagicMock()
        elem.count.return_value = 0
        elem.is_visible.return_value = False
        elem.first = elem
        return elem

    mock_page.locator.side_effect = loc_side_effect

    with pytest.raises(ValueError, match="Error de PRIZMA al guardar el test"):
        automator.guardar_creacion_test(mock_page, "Test con Error")


def test_estado_prizma_json_returns_dashboard_metrics_and_detail():
    from fastapi.testclient import TestClient
    from test_module.web import _AUTOMATION_JOBS, _AUTOMATION_LOCK, _JOBS
    from main import app, _firmar_sesion

    client = TestClient(app)
    client.cookies.set("ap_sesion", _firmar_sesion("adminlocal"))

    with _AUTOMATION_LOCK:
        _AUTOMATION_JOBS["job-kpi-test"] = {
            "owner": "adminlocal",
            "simulacion": False,
            "estado": "en_ejecucion",
            "mensaje": "Procesando tests...",
            "total": 3,
            "procesados": 1,
            "exitosos": 1,
            "pendientes": 2,
            "fallidos": 0,
            "progreso_porcentaje": 33,
            "detalle_tests": [
                {"numero": 1, "titulo": "Test 1", "preguntas": 5, "estado": "exitoso", "mensaje": "Guardado"},
                {"numero": 2, "titulo": "Test 2", "preguntas": 10, "estado": "en_proceso", "mensaje": "Procesando"},
                {"numero": 3, "titulo": "Test 3", "preguntas": 8, "estado": "pendiente", "mensaje": "En cola"},
            ],
            "logs": [],
            "captura": None,
            "error": "",
        }

    _JOBS["job-kpi-test"] = {"owner": "adminlocal", "document": None}

    response = client.get("/tests/job-kpi-test/estado-prizma/json")
    assert response.status_code == 200
    data = response.json()
    assert data["total"] == 3
    assert data["exitosos"] == 1
    assert data["pendientes"] == 2
    assert data["progreso_porcentaje"] == 33
    assert len(data["detalle_tests"]) == 3
    assert data["detalle_tests"][0]["estado"] == "exitoso"


def test_ejecutar_prizma_endpoint_respects_unchecked_simulation():
    from fastapi.testclient import TestClient
    from test_module.models import TestDocument, DocumentTest
    from test_module.web import _AUTOMATION_JOBS, _AUTOMATION_LOCK, _JOBS
    from main import app, _firmar_sesion

    client = TestClient(app)
    client.cookies.set("ap_sesion", _firmar_sesion("adminlocal"))

    doc = TestDocument(tests=[DocumentTest(title="Test 1")])
    _JOBS["job-test-real"] = {"owner": "adminlocal", "document": doc}

    # Posting WITHOUT solo_simulacion (simulating unchecked checkbox in HTML form)
    response = client.post(
        "/tests/job-test-real/ejecutar-prizma",
        data={
            "prizma_usuario": "12345",
            "prizma_clave": "secret",
            "cascada_facultad": "Fac",
            # solo_simulacion omitted!
        },
        follow_redirects=False,
    )
    assert response.status_code == 303
    with _AUTOMATION_LOCK:
        job_state = _AUTOMATION_JOBS.get("job-test-real")
        assert job_state is not None
        assert job_state["simulacion"] is False


def test_status_prizma_renders_error_banner_on_failure():
    from fastapi.testclient import TestClient
    from test_module.web import _AUTOMATION_JOBS, _AUTOMATION_LOCK, _JOBS
    from main import app, _firmar_sesion

    client = TestClient(app)
    client.cookies.set("ap_sesion", _firmar_sesion("adminlocal"))

    with _AUTOMATION_LOCK:
        _AUTOMATION_JOBS["job-error-test"] = {
            "owner": "adminlocal",
            "simulacion": False,
            "estado": "error",
            "mensaje": "Fallo en la ejecución: Selector no encontrado",
            "total": 1,
            "procesados": 0,
            "exitosos": 0,
            "pendientes": 1,
            "fallidos": 1,
            "progreso_porcentaje": 0,
            "detalle_tests": [],
            "logs": [],
            "captura": None,
            "error": "Error al presionar botón Crear",
        }

    _JOBS["job-error-test"] = {"owner": "adminlocal", "document": None}

    response = client.get("/tests/job-error-test/estado-prizma")
    assert response.status_code == 200
    assert "banner-error-superior" in response.text
    assert "La ejecución se detuvo por un error" in response.text
    assert "Error al presionar botón Crear" in response.text


def test_simplificar_mensaje_error():
    from test_module.web import simplificar_mensaje_error

    raw_pw = (
        'Error durante la automatización: Locator.click: Timeout 10000ms exceeded.\n'
        'Call log:\n'
        '- waiting for locator("button[type=\\"submit\\"]:has-text(\\"Crear\\")").first\n'
        '- locator resolved to <button ... disabled>Crear</button>\n'
        '- attempting click action 2 × waiting for element to be visible, enabled and stable - element is not enabled'
    )
    clean = simplificar_mensaje_error(raw_pw)
    assert "Call log" not in clean
    assert "Locator.click" not in clean
    assert "deshabilitado" in clean

    clean_fac = simplificar_mensaje_error("RuntimeError: No se encontró el selector 'Facultad' en el formulario de PRIZMA")
    assert "No se encontró el selector 'Facultad'" in clean_fac

    clean_crear = simplificar_mensaje_error("ValueError: No se encontró el botón 'Crear' en el formulario activo para persistir el test en PRIZMA.")
    assert "No se localizó el botón 'Crear' al lado de 'Cancelar'" in clean_crear


def test_cancelar_prizma_endpoint_registers_cancellation():
    from fastapi.testclient import TestClient
    from test_module.web import _AUTOMATION_JOBS, _AUTOMATION_LOCK, _JOBS
    from main import app, _firmar_sesion

    client = TestClient(app)
    client.cookies.set("ap_sesion", _firmar_sesion("adminlocal"))

    with _AUTOMATION_LOCK:
        _AUTOMATION_JOBS["job-stop-test"] = {
            "owner": "adminlocal",
            "simulacion": False,
            "cancelado": False,
            "estado": "en_ejecucion",
            "mensaje": "Procesando test 1...",
            "total": 3,
            "procesados": 0,
            "exitosos": 0,
            "pendientes": 3,
            "fallidos": 0,
            "progreso_porcentaje": 0,
            "detalle_tests": [],
            "logs": [],
            "captura": None,
            "error": "",
        }

    _JOBS["job-stop-test"] = {"owner": "adminlocal", "document": None}

    resp = client.post("/tests/job-stop-test/cancelar-prizma")
    assert resp.status_code == 200
    assert resp.json()["ok"] is True

    with _AUTOMATION_LOCK:
        state = _AUTOMATION_JOBS["job-stop-test"]
        assert state["cancelado"] is True


def test_status_prizma_renders_stop_button_and_detenido_banner():
    from fastapi.testclient import TestClient
    from test_module.web import _AUTOMATION_JOBS, _AUTOMATION_LOCK, _JOBS
    from main import app, _firmar_sesion

    client = TestClient(app)
    client.cookies.set("ap_sesion", _firmar_sesion("adminlocal"))

    # When in execution: Stop button should be visible (display: inline-flex)
    with _AUTOMATION_LOCK:
        _AUTOMATION_JOBS["job-active"] = {
            "owner": "adminlocal",
            "simulacion": False,
            "estado": "en_ejecucion",
            "mensaje": "Creando test...",
            "total": 2,
            "procesados": 0,
            "exitosos": 0,
            "pendientes": 2,
            "fallidos": 0,
            "progreso_porcentaje": 0,
            "detalle_tests": [],
            "logs": [],
            "captura": None,
            "error": "",
        }
    _JOBS["job-active"] = {"owner": "adminlocal", "document": None}
    r_act = client.get("/tests/job-active/estado-prizma")
    assert r_act.status_code == 200
    assert "btn-detener-auto" in r_act.text
    assert "display:inline-flex" in r_act.text

    # When detenido: Detenido banner should be displayed
    with _AUTOMATION_LOCK:
        _AUTOMATION_JOBS["job-detenido"] = {
            "owner": "adminlocal",
            "simulacion": False,
            "estado": "detenido",
            "mensaje": "Automatización detenida por el usuario.",
            "total": 2,
            "procesados": 1,
            "exitosos": 1,
            "pendientes": 1,
            "fallidos": 0,
            "progreso_porcentaje": 50,
            "detalle_tests": [],
            "logs": [],
            "captura": None,
            "error": "",
        }
    _JOBS["job-detenido"] = {"owner": "adminlocal", "document": None}
    r_det = client.get("/tests/job-detenido/estado-prizma")
    assert r_det.status_code == 200
    assert "banner-detenido-superior" in r_det.text
    assert "Automatización detenida por el usuario" in r_det.text


def test_guardar_creacion_test_locates_crear_beside_cancelar_and_waits():
    automator = PrizmaTestAutomator()
    automator._tomar_captura = MagicMock(return_value="")

    mock_page = MagicMock()

    mock_btn_crear = MagicMock()
    mock_btn_crear.is_visible.return_value = True
    # Start disabled, then becomes enabled
    mock_btn_crear.is_enabled.side_effect = [False, False, True]
    mock_btn_crear.inner_text.return_value = "Crear"

    mock_btn_cancelar = MagicMock()
    mock_btn_cancelar.is_visible.return_value = True
    mock_btn_cancelar.inner_text.return_value = "Cancelar"

    # Set up sibling xpath locator
    mock_hermanos = MagicMock()
    mock_hermanos.count.return_value = 1
    mock_hermanos.nth.return_value = mock_btn_crear
    mock_btn_cancelar.locator.return_value = mock_hermanos

    mock_campo_nombre = MagicMock()
    mock_campo_nombre.count.return_value = 0
    mock_campo_nombre.is_visible.return_value = False

    def loc_side_effect(sel):
        if "Cancelar" in sel and ("~" in sel or "+" in sel):
            elem = MagicMock()
            elem.is_visible.return_value = False
            return elem
        if "Cancelar" in sel:
            mock_loc = MagicMock()
            mock_loc.filter.return_value.first = mock_btn_cancelar
            mock_loc.first = mock_btn_cancelar
            return mock_loc
        if "Nombre del test" in sel:
            return mock_campo_nombre
        elem = MagicMock()
        elem.count.return_value = 0
        elem.is_visible.return_value = False
        elem.first = elem
        return elem

    mock_page.locator.side_effect = loc_side_effect

    res = automator.guardar_creacion_test(mock_page, "Test Espera")
    assert res is True
    mock_btn_crear.click.assert_called_once()


def test_ejecutar_cargue_honors_cancellation():
    from test_module.models import DocumentTest

    automator = PrizmaTestAutomator()
    automator._tomar_captura = MagicMock(return_value="")

    with patch("test_module.automation.sync_playwright") as mock_pw:
        mock_p_ctx = MagicMock()
        mock_pw.return_value.__enter__.return_value = mock_p_ctx
        mock_browser = MagicMock()
        mock_p_ctx.chromium.launch.return_value = mock_browser
        mock_context = MagicMock()
        mock_browser.new_context.return_value = mock_context
        mock_page = MagicMock()
        mock_context.new_page.return_value = mock_page

        automator.login = MagicMock(return_value=True)
        automator.navegar_a_crear_test = MagicMock()
        automator.seleccionar_mui = MagicMock()
        automator.llenar_descripcion_test = MagicMock()
        automator.inyectar_json_test = MagicMock(return_value=True)

        tests = [
            DocumentTest(title="Test 1"),
            DocumentTest(title="Test 2"),
            DocumentTest(title="Test 3"),
        ]

        # Stop after test 1
        processed_count = 0
        def cancel_check():
            return processed_count >= 1

        callbacks = []
        def test_cb(idx, status, msg, err):
            callbacks.append((idx, status))
            if status == "exitoso":
                nonlocal processed_count
                processed_count += 1

        res = automator.ejecutar_cargue(
            usuario_prizma="user",
            clave_prizma="pass",
            tests=tests,
            cascada={},
            solo_simulacion=True,
            test_callback=test_cb,
            cancel_checker=cancel_check,
        )

        assert res["ok"] is True
        assert res["detenido"] is True
        assert res["procesados"] == 1
        # Tests 2 and 3 should have been marked cancelado
        assert (2, "cancelado") in callbacks
        assert (3, "cancelado") in callbacks


def test_extraer_valor_campo_formulario_input_y_container():
    automator = PrizmaTestAutomator()
    mock_page = MagicMock()

    mock_input = MagicMock()
    mock_input.input_value.return_value = "Facultad de Ingeniería"
    mock_input.is_visible.return_value = True

    automator.localizar_campo_formulario = MagicMock(return_value=mock_input)
    val = automator.extraer_valor_campo_formulario(mock_page, "Facultad")
    assert val == "Facultad de Ingeniería"


def test_ejecutar_test_job_exitoso():
    from test_module.web import ejecutar_test_job, _JOBS, _JOBS_LOCK, _AUTOMATION_JOBS
    from test_module.models import TestDocument, DocumentTest

    job_id = "test-job-exec-1"
    doc = TestDocument(tests=[DocumentTest(title="Parcial 1", week="1", cut="1", description="", sections=[])])
    with _JOBS_LOCK:
        _JOBS[job_id] = {
            "owner": "adminlocal",
            "document": doc,
            "cascada_datos": {"programa": "Sistemas", "pensum": "P1", "nivel": "1", "asignatura": "Algoritmos"},
        }

    with patch("test_module.web.PrizmaTestAutomator.ejecutar_cargue") as mock_run:
        mock_run.return_value = {"ok": True, "exitosos": 1, "procesados": 1}
        res = ejecutar_test_job(
            job_id=job_id,
            usuario_prizma="user123",
            clave_prizma="pass123",
            headless=True,
        )
        assert res["ok"] is True
        assert _AUTOMATION_JOBS[job_id]["estado"] == "completado"
        assert _AUTOMATION_JOBS[job_id]["progreso_porcentaje"] == 100


def test_ejecutar_test_en_cargue_integracion():
    from main import _ejecutar_test_en_cargue
    from test_module.web import _JOBS, _JOBS_LOCK
    from test_module.models import TestDocument, DocumentTest

    job_id = "test-cargue-integ-1"
    doc = TestDocument(tests=[DocumentTest(title="Evaluación Final", week="16", cut="3", description="", sections=[])])
    with _JOBS_LOCK:
        _JOBS[job_id] = {
            "owner": "adminlocal",
            "document": doc,
            "cascada_datos": {"programa": "Derecho", "pensum": "P2", "nivel": "2", "asignatura": "Constitucional"},
        }

    trabajo = {
        "detalle_actividades": [
            {"numero": 1, "nombre": "OVI: Intro", "estado": "ok", "error": ""}
        ],
        "total": 1,
        "procesadas": 1,
        "exitosas": 1,
        "errores": 0,
        "terminado": False,
    }

    with patch("main.ejecutar_test_job", create=True) as mock_exec, patch("test_module.web.ejecutar_test_job") as mock_exec_w:
        def fake_exec(*args, **kwargs):
            cb = kwargs.get("on_test_callback")
            if cb:
                cb(1, "exitoso", "Guardado OK")
            return {"ok": True, "exitosos": 1, "procesados": 1}

        mock_exec.side_effect = fake_exec
        mock_exec_w.side_effect = fake_exec

        _ejecutar_test_en_cargue(trabajo, job_id, "user", "pass")

        assert len(trabajo["detalle_actividades"]) == 2
        test_act = trabajo["detalle_actividades"][1]
        assert "Test Evaluativo" in test_act["nombre"]
        assert test_act["estado"] == "ok"
        assert trabajo["terminado"] is True
        assert trabajo["etapa"] == "finalizado"
        assert trabajo["exitosas"] == 2
        assert trabajo["total"] == 2


def test_endpoint_estado_activo_json_y_home_card():
    from fastapi.testclient import TestClient
    from test_module.web import router, _AUTOMATION_JOBS, _AUTOMATION_LOCK
    from fastapi import FastAPI

    app = FastAPI()
    app.include_router(router)
    client = TestClient(app)

    # 1. Sin trabajo activo
    with _AUTOMATION_LOCK:
        _AUTOMATION_JOBS.clear()

    resp_empty = client.get("/tests/estado-activo/json")
    assert resp_empty.status_code == 200
    assert resp_empty.json()["activo"] is False

    # 2. Con trabajo activo
    with _AUTOMATION_LOCK:
        _AUTOMATION_JOBS["job-activo-xyz"] = {
            "owner": "test_user",
            "estado": "en_ejecucion",
            "mensaje": "Cargando pregunta 2 de 5...",
            "progreso_porcentaje": 40,
            "total": 1,
            "procesados": 0,
            "exitosos": 0,
            "fallidos": 0,
            "asignatura": "Cálculo I",
            "titulo_doc": "Parcial Cálculo.docx",
        }

    resp_active = client.get("/tests/estado-activo/json")
    assert resp_active.status_code == 200
    data = resp_active.json()
    assert data["activo"] is True
    assert data["job_id"] == "job-activo-xyz"
    assert data["progreso_porcentaje"] == 40

    # 3. Verificar que /tests renderiza la tarjeta
    resp_home = client.get("/tests")
    assert resp_home.status_code == 200
    assert "card-estado-test-activo" in resp_home.text
    assert "Cálculo I" in resp_home.text
