from io import BytesIO

from docx import Document as DocxDocument

from test_module.models import Severity
from test_module.parser import parse_document_text, parse_docx_bytes
from test_module.validation import validate_document
from test_module.serializer import serialize_test


def test_parse_document_supports_all_prizma_question_types_and_multiple_answers():
    text = """Semana 1 - Corte 1
Título del TEST
Evaluación de fundamentos
Estimados estudiantes.
Responde con base en los contenidos estudiados.
Verdadero/Falso
1. El agua hierve a 100 °C. [R: Verdadero]
Preguntas de única respuesta
1. ¿Cuál es la capital de Colombia?
a) Cali
b) Bogotá
c) Medellín
[R: b]
2. Selecciona los elementos correctos.
a) A
b) B
c) C
[R: a,c]
Emparejamiento
1. Relaciona cada concepto.
Columna A
1. Python
2. HTML
Columna B
a) Lenguaje de programación
b) Lenguaje de marcado
[R: 1-a, 2-b]
Pregunta abierta
1. Explica qué es una variable.
"""

    document = parse_document_text(text)

    assert len(document.tests) == 1
    test = document.tests[0]
    assert test.title == "Evaluación de fundamentos"
    assert test.week == "1"
    assert test.cut == "1"
    assert [section.code for section in test.sections] == [
        "trueOrFalse",
        "multipleChoice",
        "matching",
        "writeAnswer",
    ]
    assert test.sections[1].questions[1].correct_option_letters == ["a", "c"]
    assert len(test.sections[2].questions[0].pairs) == 2
    assert test.sections[3].questions[0].prompt == "Explica qué es una variable."


def test_validation_marks_missing_closed_answer_as_error():
    document = parse_document_text(
        """Título del TEST
        Diagnóstico
        Selección múltiple
        1. ¿Cuál opción es correcta?
        a) Primera
        b) Segunda
        """
    )

    issues = validate_document(document)

    assert any(
        issue.severity is Severity.ERROR
        and "respuesta" in issue.message.lower()
        for issue in issues
    )


def test_docx_parser_reads_paragraphs_and_tables():
    source = DocxDocument()
    source.add_paragraph("Título del TEST")
    source.add_paragraph("Prueba desde Word")
    source.add_paragraph("Verdadero/Falso")
    source.add_paragraph("1. La prueba funciona. [R: Verdadero]")
    table = source.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Semana 3"
    table.rows[0].cells[1].text = "Corte 2"
    buffer = BytesIO()
    source.save(buffer)

    document = parse_docx_bytes(buffer.getvalue())

    assert document.tests[0].title == "Prueba desde Word"
    assert document.tests[0].week == "3"
    assert document.tests[0].cut == "2"
    assert document.tests[0].sections[0].questions[0].answer is True


def test_serializer_emits_prizma_import_shape_with_uuid_and_html():
    document = parse_document_text(
        """Título del TEST
        Prueba JSON
        Verdadero/Falso
        1. La tierra es redonda. [R: Verdadero]
        """
    )

    payload = serialize_test(document.tests[0])

    assert payload["test_information"]["title"] == "Prueba JSON"
    activity = payload["test_information"]["activities"][0]
    assert activity["activity_type"]["code"] == "trueOrFalse"
    assert activity["activity_type"]["id"] == "122d0f18-7aa9-44e1-9dc1-2c0851a10940"
    assert activity["contents"][0]["content"] == "<p>La tierra es redonda.</p>"
    assert activity["contents"][0]["content_options"][0]["right_answer"] is True


def test_docx_parser_splits_multiple_tests_and_extracts_titles_from_tables():
    source = DocxDocument()
    table = source.add_table(rows=6, cols=2)
    # Test 1
    table.rows[0].cells[0].text = "Semana/corte"
    table.rows[0].cells[1].text = "Semana 3 - corte 1"
    table.rows[1].cells[0].text = "Título del TEST"
    table.rows[1].cells[1].text = "La Infraestructura o Red de Transporte."
    table.rows[2].cells[0].text = "Verdadero/Falso\n1. Pregunta uno. [R: Verdadero]"
    table.rows[2].cells[1].text = "Verdadero/Falso\n1. Pregunta uno. [R: Verdadero]"
    # Test 2
    table.rows[3].cells[0].text = "Semana/corte"
    table.rows[3].cells[1].text = "Semana 5 - corte 2"
    table.rows[4].cells[0].text = "Título del TEST"
    table.rows[4].cells[1].text = "Los Flujos de Transporte"
    table.rows[5].cells[0].text = "Verdadero/Falso\n1. Pregunta dos. [R: Falso]"
    table.rows[5].cells[1].text = "Verdadero/Falso\n1. Pregunta dos. [R: Falso]"

    buffer = BytesIO()
    source.save(buffer)
    document = parse_docx_bytes(buffer.getvalue())

    assert len(document.tests) == 2
    assert document.tests[0].title == "La Infraestructura o Red de Transporte"
    assert document.tests[0].week == "3"
    assert document.tests[0].cut == "1"
    assert document.tests[1].title == "Los Flujos de Transporte"
    assert document.tests[1].week == "5"
    assert document.tests[1].cut == "2"


def test_serializer_emits_matching_pairs_in_contents():
    document = parse_document_text(
        """Título del TEST
        Prueba Emparejamiento
        Emparejamiento
        1. Relaciona cada concepto.
        Columna A
        1. Python
        2. HTML
        Columna B
        a) Lenguaje de programación
        b) Lenguaje de marcado
        [R: 1-a, 2-b]
        """
    )
    payload = serialize_test(document.tests[0])
    activity = payload["test_information"]["activities"][0]
    assert activity["activity_type"]["code"] == "matching"
    assert len(activity["contents"]) == 2
    assert activity["contents"][0]["content"] == "Python"
    assert activity["contents"][0]["content_options"][0]["text"] == "Lenguaje de programación"
    assert activity["contents"][0]["content_options"][0]["right_answer"] is True


def test_serializer_splits_multiple_matching_questions_into_separate_activities():
    document = parse_document_text(
        """Título del TEST
        Prueba Multiple Emparejamiento
        Emparejamiento
        1. Relaciona países con sus capitales.
        Columna A
        1. Colombia
        2. Francia
        Columna B
        a) Bogotá
        b) París
        [R: 1-a, 2-b]
        2. Relaciona monedas con países.
        Columna A
        1. Peso
        2. Euro
        Columna B
        a) Colombia
        b) Francia
        [R: 1-a, 2-b]
        3. Relaciona autores con obras.
        Columna A
        1. García Márquez
        2. Cervantes
        Columna B
        a) Cien años de soledad
        b) Don Quijote
        [R: 1-a, 2-b]
        """
    )
    payload = serialize_test(document.tests[0])
    activities = payload["test_information"]["activities"]
    assert len(activities) == 3
    assert activities[0]["name_section"] == "Emparejamiento 1"
    assert "países con sus capitales" in activities[0]["description_section"]
    assert activities[0]["contents"][0]["content"] == "Colombia"
    assert activities[0]["contents"][0]["content_options"][0]["text"] == "Bogotá"

    assert activities[1]["name_section"] == "Emparejamiento 2"
    assert "monedas con países" in activities[1]["description_section"]
    assert activities[1]["contents"][0]["content"] == "Peso"

    assert activities[2]["name_section"] == "Emparejamiento 3"
    assert "autores con obras" in activities[2]["description_section"]
    assert activities[2]["contents"][0]["content"] == "García Márquez"


def test_parse_true_false_strips_options_from_prompt():
    text = """Título del TEST
    Test Evaluación
    Verdadero/Falso
    1. La fase de despacho terrestre en los puertos se refiere a la llegada.
    ( ) Verdadero     ( ) Falso
    [R: Falso]
    2. Otra pregunta con opciones en la misma línea. ( ) Verdadero ( ) Falso [R: Verdadero]
    """
    document = parse_document_text(text)
    tf_section = document.tests[0].sections[0]
    assert len(tf_section.questions) == 2
    assert tf_section.questions[0].prompt == "La fase de despacho terrestre en los puertos se refiere a la llegada."
    assert tf_section.questions[0].answer is False
    assert tf_section.questions[1].prompt == "Otra pregunta con opciones en la misma línea."
    assert tf_section.questions[1].answer is True


def test_parse_docx_bytes_extracts_program_only_with_code():
    doc_with_code = DocxDocument()
    table = doc_with_code.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Programa"
    table.cell(0, 1).text = "COD111 - Negocios Internacionales"
    table.cell(1, 0).text = "Título del TEST"
    table.cell(1, 1).text = "Examen con Código"
    p = doc_with_code.add_paragraph("Verdadero/Falso\n1. Pregunta. [R: Verdadero]")
    buf1 = BytesIO()
    doc_with_code.save(buf1)

    parsed_code = parse_docx_bytes(buf1.getvalue())
    assert parsed_code.program == "COD111 - Negocios Internacionales"

    # Test without code: should remain blank
    doc_no_code = DocxDocument()
    table2 = doc_no_code.add_table(rows=2, cols=2)
    table2.cell(0, 0).text = "Programa"
    table2.cell(0, 1).text = "Negocios Internacionales"
    table2.cell(1, 0).text = "Título del TEST"
    table2.cell(1, 1).text = "Examen sin Código"
    doc_no_code.add_paragraph("Verdadero/Falso\n1. Pregunta. [R: Verdadero]")
    buf2 = BytesIO()
    doc_no_code.save(buf2)

    parsed_no_code = parse_docx_bytes(buf2.getvalue())
    assert parsed_no_code.program == ""


